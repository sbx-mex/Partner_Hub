import unittest
import json
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from engine.partner_engine.generator import (
    RUNTIME_HEADERS, build_navigation, build_partner_data, ceco_from_ccosto, directory_ceco,
    employee_id, is_active_status, iso_date,
)
from tools.audit_sources import audit_sources


class PartnerEngineTests(unittest.TestCase):
    def test_ccosto_uses_last_five_digits_only(self):
        self.assertEqual(ceco_from_ccosto("01938262"), "38262")
        self.assertNotEqual(ceco_from_ccosto("01938262"), "38362")
        self.assertEqual(ceco_from_ccosto("01538656"), "38656")
        self.assertEqual(ceco_from_ccosto(1538656), "38656")
        self.assertEqual(ceco_from_ccosto("00042"), "00042")
        self.assertEqual(ceco_from_ccosto("15A38656"), "")
        self.assertEqual(directory_ceco(42), "00042")
        self.assertEqual(directory_ceco("00042"), "00042")

    def test_employee_id_preserves_leading_zero(self):
        self.assertEqual(employee_id(33903), "033903")
        self.assertEqual(employee_id("033903"), "033903")

    def test_date_serialization(self):
        self.assertEqual(iso_date(datetime(2026, 7, 31)), "2026-07-31")

    def test_active_status_does_not_accept_inactive(self):
        self.assertTrue(is_active_status("ACTIVOS"))
        self.assertFalse(is_active_status("INACTIVOS"))
        self.assertFalse(is_active_status("BAJAS"))

    def test_navigation_is_grouped_and_sorted(self):
        def row(region, dm, store):
            values = [""] * len(RUNTIME_HEADERS)
            values[RUNTIME_HEADERS.index("REGION")] = region
            values[RUNTIME_HEADERS.index("DM")] = dm
            values[RUNTIME_HEADERS.index("NOM_CCOSTO")] = store
            return values
        nav = build_navigation([row("NORTE", "DM B", "Tienda 2"), row("NORTE", "DM A", "Tienda 1")], RUNTIME_HEADERS)
        self.assertEqual(nav[0]["name"], "NORTE")
        self.assertEqual([item["name"] for item in nav[0]["dms"]], ["DM A", "DM B"])


class WorkbookIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.root = Path(self.folder.name)
        self.query = self.root / "Query.xlsx"
        self.directory = self.root / "Directorio.xlsx"
        self.output = self.root / "partners.js"
        self.audit = self.root / "audit.json"
        self.query_headers = ["NOM_PUESTO", "CCOSTO", "F.NAC", "NUM_EMP", "NOMBRE", "F_BAJA", "TURNO", "F_INGRESO", "NOM_CCOSTO"]
        self.query_rows = [{"NOM_PUESTO": "Barista", "CCOSTO": 1538656, "F.NAC": date(1999, 2, 18),
                            "NUM_EMP": 123, "NOMBRE": "Ana <script>", "F_BAJA": None,
                            "TURNO": "Matutino", "F_INGRESO": date(2020, 8, 1), "NOM_CCOSTO": "Nombre anterior"}]
        self.directory_headers = ["DM", "Estatus", "Región", "CC Nombre", "CC"]
        self.directory_rows = [{"CC": 38656, "CC Nombre": "Tienda Centro", "Región": "Centro Norte", "Estatus": "Abierta", "DM": "Diana"}]

    def save(self):
        for path, sheet_name, headers, rows in ((self.query, "query", self.query_headers, self.query_rows),
                                               (self.directory, "Directorio", self.directory_headers, self.directory_rows)):
            book = Workbook()
            sheet = book.active
            sheet.title = sheet_name
            sheet.append(headers)
            if path == self.query:
                sheet["L1"] = "Corte de información"
                sheet["M1"] = date(2026, 9, 28)
            for row in rows:
                sheet.append([row.get(header) for header in headers])
            book.save(path)

    def build(self):
        self.save()
        return build_partner_data(self.query, self.output, self.audit,
                                  directory_path=self.directory, as_of=date(2026, 9, 28))

    def test_reordered_headers_join_and_minimal_public_fields(self):
        result = self.build()
        self.assertEqual((result.published_rows, result.stores, result.regions), (1, 1, 1))
        data = self.output.read_text(encoding="utf-8")
        rows = json.loads(data.split("window.PARTNER_ROWS=", 1)[1].rstrip(";\n"))
        self.assertEqual(rows[0][RUNTIME_HEADERS.index("CECO")], "38656")
        self.assertEqual(rows[0][RUNTIME_HEADERS.index("NOM_CCOSTO")], "Tienda Centro")
        self.assertEqual(rows[0][RUNTIME_HEADERS.index("CUMPLE_MMDD")], "--02-18")
        self.assertNotIn("1999", data)
        self.assertNotIn("000123", data)
        audit = json.loads(self.audit.read_text(encoding="utf-8"))
        self.assertEqual(audit["queryStoreNameDifferences"], 1)
        self.assertEqual(audit["periodEnd"], "2026-09-28")

    def test_identical_duplicate_and_terminated_are_excluded(self):
        self.query_rows.append(dict(self.query_rows[0]))
        self.query_rows.append({**self.query_rows[0], "NUM_EMP": 124, "F_BAJA": date(2026, 9, 1)})
        self.build()
        audit = json.loads(self.audit.read_text(encoding="utf-8"))
        self.assertEqual((audit["sourceRows"], audit["publishedRows"], audit["terminatedExcluded"],
                          audit["identicalEmployeeRowsRemoved"]), (3, 1, 1, 1))

    def test_invalid_join_never_overwrites_prior_output(self):
        self.output.write_text("anterior", encoding="utf-8")
        self.audit.write_text("anterior", encoding="utf-8")
        self.query_rows[0]["CCOSTO"] = 1538657
        with self.assertRaisesRegex(ValueError, "no existe"):
            self.build()
        self.assertEqual(self.output.read_text(encoding="utf-8"), "anterior")
        self.assertEqual(self.audit.read_text(encoding="utf-8"), "anterior")

    def test_closed_or_conflicting_directory_rejected(self):
        self.directory_rows[0]["Estatus"] = "Cerrada"
        with self.assertRaisesRegex(ValueError, "no está abierto"):
            self.build()
        self.directory_rows[0]["Estatus"] = "Abierta"
        self.directory_rows.append({**self.directory_rows[0], "DM": "Otra DM"})
        with self.assertRaisesRegex(ValueError, "contradictorio"):
            self.build()

    def test_directory_numeric_ceco_with_lost_leading_zero(self):
        self.directory_rows[0]["CC"] = 42
        self.query_rows[0]["CCOSTO"] = "0100042"
        self.build()
        self.assertIn('"00042"', self.output.read_text(encoding="utf-8"))

    def test_duplicate_headers_rejected(self):
        self.query_headers.append("NOMBRE")
        with self.assertRaisesRegex(ValueError, "Encabezado duplicado"):
            self.build()

    def test_cutoff_is_required_and_does_not_follow_execution_date(self):
        self.save()
        from openpyxl import load_workbook
        book = load_workbook(self.query)
        book.active["M1"] = None
        book.save(self.query)
        with self.assertRaisesRegex(ValueError, "M1"):
            build_partner_data(self.query, self.output, self.audit, directory_path=self.directory)
        self.assertFalse(self.output.exists())
        self.save()
        with self.assertRaisesRegex(ValueError, "no coincide"):
            build_partner_data(self.query, self.output, self.audit, directory_path=self.directory,
                               as_of=date(2026, 9, 29))

    def test_name_matching_another_ceco_blocks_wrong_dm(self):
        self.directory_rows.append({"CC": 38362, "CC Nombre": "Nombre anterior", "Región": "Otra",
                                    "Estatus": "Abierta", "DM": "Otro DM"})
        with self.assertRaisesRegex(ValueError, "CeCos distintos"):
            self.build()

    def test_future_termination_remains_active_at_cutoff(self):
        self.query_rows[0]["F_BAJA"] = date(2026, 10, 1)
        result = self.build()
        self.assertEqual(result.published_rows, 1)

    def test_preflight_reports_missing_ceco_without_employee_details(self):
        self.query_rows.append({**self.query_rows[0], "CCOSTO": "01938362", "NUM_EMP": 555,
                                "NOMBRE": "Persona privada"})
        self.save()
        report = audit_sources(self.query, self.directory)
        self.assertFalse(report["canPublish"])
        self.assertEqual(report["missingCeCos"], [{"ceco": "38362", "rows": 1}])
        self.assertNotIn("Persona privada", json.dumps(report, ensure_ascii=False))

    def test_output_cannot_replace_source_book(self):
        self.save()
        original = self.query.read_bytes()
        with self.assertRaisesRegex(ValueError, "rutas diferentes"):
            build_partner_data(self.query, self.query, self.audit, directory_path=self.directory)
        self.assertEqual(self.query.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
