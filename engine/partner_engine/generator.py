"""Carga validada de Partner Hub desde Query y Directorio."""
from __future__ import annotations

import hashlib
import json
import math
import os
import re
import tempfile
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable
from openpyxl import load_workbook

EXPECTED_HEADERS = ("NUM_EMP", "NOMBRE", "F_INGRESO", "F_BAJA", "CCOSTO", "TURNO", "NOM_PUESTO", "F.NAC")
DIRECTORY_HEADERS = ("CC", "CC Nombre", "Región", "Estatus", "DM")
# Campos mínimos públicos: sin número de empleado, año de nacimiento ni CCOSTO original.
RUNTIME_HEADERS = ("NOMBRE", "F_INGRESO", "CECO", "NOM_CCOSTO", "TURNO", "NOM_PUESTO", "REGION", "DM", "CUMPLE_MMDD")
MONTHS = ("Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre")

@dataclass(frozen=True)
class BuildResult:
    source_rows: int
    published_rows: int
    period_end: str
    month: int
    year: int
    regions: int
    dms: int
    stores: int
    output: str
    audit_output: str

def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip() if value is not None else ""

def sort_key(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", clean_text(value).casefold())
    return "".join(ch for ch in normalized if not unicodedata.combining(ch))

def header_key(value: Any) -> str:
    return sort_key(value)

def employee_id(value: Any) -> str:
    if isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or not float(value).is_integer():
            return ""
        value = str(int(value))
    text = clean_text(value)
    return text.split(".", 1)[0].zfill(6) if re.fullmatch(r"\d+(?:\.0+)?", text) else ""

def ceco_from_ccosto(value: Any) -> str:
    """Últimos cinco dígitos; conserva ceros si Excel entrega texto."""
    if isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or not float(value).is_integer():
            return ""
        value = str(int(value))
    text = clean_text(value)
    return text.split(".", 1)[0][-5:] if re.fullmatch(r"\d{5,12}(?:\.0+)?", text) else ""

def directory_ceco(value: Any) -> str:
    """Excel puede guardar el CeCo con ceros iniciales como número."""
    if isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        return str(int(value)).zfill(5) if 0 <= value <= 99999 and math.isfinite(value) and float(value).is_integer() else ""
    text = clean_text(value)
    return text.split(".", 1)[0] if re.fullmatch(r"\d{5}(?:\.0+)?", text) else ""

def is_active_status(value: Any) -> bool:
    return sort_key(value).startswith("activo")

def iso_date(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = clean_text(value)
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError("Fecha no reconocida en Query")

def query_cutoff(sheet: Any) -> date:
    """M1 es el corte explícito del Query; nunca se sustituye por la fecha del build."""
    label = sort_key(sheet["L1"].value)
    if "corte" not in label:
        raise ValueError("Query L1 debe indicar 'Corte de información'")
    try:
        value = iso_date(sheet["M1"].value)
        cutoff = date.fromisoformat(value)
    except (ValueError, TypeError) as error:
        raise ValueError("Query M1 requiere una fecha de corte válida") from error
    if not value or cutoff > date.today():
        raise ValueError("Query M1 requiere una fecha de corte válida, no futura")
    return cutoff

def find_header_row(sheet: Any, expected: Iterable[str], max_rows: int = 10) -> tuple[int, list[str]]:
    required = {header_key(name) for name in expected}
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=max_rows, values_only=True), 1):
        headers = [clean_text(value) for value in row]
        if required.issubset({header_key(value) for value in headers if value}):
            return row_number, headers
    raise ValueError("Faltan encabezados requeridos: " + ", ".join(expected))

def column_positions(headers: list[str]) -> dict[str, int]:
    positions = {}
    for index, header in enumerate(headers):
        key = header_key(header)
        if key:
            if key in positions:
                raise ValueError(f"Encabezado duplicado o equivalente: {header}")
            positions[key] = index
    return positions

def cell(row: tuple[Any, ...], columns: dict[str, int], header: str) -> Any:
    index = columns[header_key(header)]
    return row[index] if index < len(row) else None

def build_navigation(rows: list[list[Any]], headers: Iterable[str]) -> list[dict[str, Any]]:
    positions = {name: index for index, name in enumerate(headers)}
    groups: dict[tuple[str, str], set[str]] = defaultdict(set)
    labels = {}
    for row in rows:
        region, dm, store = (clean_text(row[positions[name]]) for name in ("REGION", "DM", "NOM_CCOSTO"))
        key = (sort_key(region), sort_key(dm))
        labels[key] = (region, dm)
        groups[key].add(store)
    regions: dict[str, list[dict[str, Any]]] = defaultdict(list)
    region_labels = {}
    for key in sorted(groups):
        region, dm = labels[key]
        region_labels[key[0]] = region
        regions[key[0]].append({"name": dm, "stores": sorted(groups[key], key=sort_key)})
    return [{"name": region_labels[key], "dms": regions[key]} for key in sorted(regions)]

def _digest(path: Path) -> str:
    sha = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            sha.update(chunk)
    return sha.hexdigest()

def _atomic_write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(content)
        os.replace(temporary, path)
    except BaseException:
        Path(temporary).unlink(missing_ok=True)
        raise

def _baseline_count(path: str | Path | None) -> int | None:
    if not path or not Path(path).is_file():
        return None
    match = re.search(r"window\.PARTNER_ROWS=(.*);\s*$", Path(path).read_text(encoding="utf-8"), re.S)
    if not match:
        return None
    try:
        return len(json.loads(match.group(1)))
    except (TypeError, ValueError):
        return None

def load_directory(path: str | Path) -> tuple[dict[str, dict[str, str]], set[str], dict[str, int]]:
    directory = Path(path)
    if not directory.is_file():
        raise FileNotFoundError(f"Falta Directorio.xlsx: {directory}")
    workbook = load_workbook(directory, read_only=True, data_only=True)
    try:
        sheet = workbook["Directorio"] if "Directorio" in workbook.sheetnames else workbook.active
        first, headers = find_header_row(sheet, DIRECTORY_HEADERS)
        columns = column_positions(headers)
        stores, conflicts, duplicates, invalid_rows = {}, set(), 0, 0
        for row in sheet.iter_rows(min_row=first + 1, values_only=True):
            if not any(value is not None for value in row):
                continue
            ceco = directory_ceco(cell(row, columns, "CC"))
            if not ceco:
                invalid_rows += 1
                continue
            record = {key: clean_text(cell(row, columns, source)) for key, source in
                      (("store", "CC Nombre"), ("region", "Región"), ("status", "Estatus"), ("dm", "DM"))}
            if not all(record.values()):
                invalid_rows += 1
                conflicts.add(ceco)
                continue
            if ceco in stores:
                duplicates += 1
                if any(sort_key(stores[ceco][key]) != sort_key(record[key]) for key in record):
                    conflicts.add(ceco)
            else:
                stores[ceco] = record
        return stores, conflicts, {"rows": len(stores) + duplicates + invalid_rows,
                                   "duplicateCeCos": duplicates, "invalidRowsExcluded": invalid_rows,
                                   "conflictingCeCosExcluded": len(conflicts)}
    finally:
        workbook.close()

def build_partner_data(
    excel_path: str | Path, output_path: str | Path, audit_path: str | Path, *,
    directory_path: str | Path = "data/Directorio.xlsx", sheet_name: str = "query",
    expected_month: int | None = None, expected_year: int | None = None,
    baseline_path: str | Path | None = None, as_of: date | None = None,
) -> BuildResult:
    excel, directory = Path(excel_path), Path(directory_path)
    output, audit_output = Path(output_path), Path(audit_path)
    if len({path.resolve() for path in (excel, directory, output, audit_output)}) != 4:
        raise ValueError("Query, Directorio, salida y auditoría requieren rutas diferentes")
    if not excel.is_file():
        raise FileNotFoundError(f"Falta Query.xlsx: {excel}")
    stores, conflicts, directory_audit = load_directory(directory)
    workbook = load_workbook(excel, read_only=True, data_only=True)
    try:
        matches = [name for name in workbook.sheetnames if header_key(name) == header_key(sheet_name)]
        if len(matches) != 1:
            raise ValueError(f"No se encontró una sola pestaña Query: {workbook.sheetnames}")
        sheet = workbook[matches[0]]
        cutoff = query_cutoff(sheet)
        if as_of is not None and as_of != cutoff:
            raise ValueError(f"--as-of {as_of} no coincide con el corte M1 {cutoff}")
        if expected_month is not None and cutoff.month != expected_month:
            raise ValueError(f"Mes esperado {expected_month}; corte M1 {cutoff}")
        if expected_year is not None and cutoff.year != expected_year:
            raise ValueError(f"Año esperado {expected_year}; corte M1 {cutoff}")
        first, headers = find_header_row(sheet, EXPECTED_HEADERS)
        columns = column_positions(headers)
        name_to_cecos: dict[str, set[str]] = defaultdict(set)
        for code, store in stores.items():
            name_to_cecos[sort_key(store["store"])].add(code)
        source_rows = terminated = name_mismatches = identical_rows = 0
        skipped: dict[str, int] = defaultdict(int)
        candidates: dict[str, tuple[Any, ...]] = {}
        invalid_employees: set[str] = set()
        for row in sheet.iter_rows(min_row=first + 1, values_only=True):
            if not any(value is not None for value in row):
                continue
            source_rows += 1
            try:
                termination = iso_date(cell(row, columns, "F_BAJA"))
            except ValueError:
                skipped["invalidRow"] += 1
                continue
            if termination and date.fromisoformat(termination) <= cutoff:
                terminated += 1
                continue
            number = employee_id(cell(row, columns, "NUM_EMP"))
            name = clean_text(cell(row, columns, "NOMBRE"))
            ceco = ceco_from_ccosto(cell(row, columns, "CCOSTO"))
            if not number or not name or not ceco:
                skipped["invalidRow"] += 1
                continue
            if ceco in conflicts:
                skipped["conflictingDirectory"] += 1
                continue
            if ceco not in stores:
                skipped["missingCeCo"] += 1
                continue
            store = stores[ceco]
            if sort_key(store["status"]) != "abierta":
                skipped["notOpen"] += 1
                continue
            query_name = sort_key(cell(row, columns, "NOM_CCOSTO")) if header_key("NOM_CCOSTO") in columns else ""
            other_cecos = name_to_cecos.get(query_name, set()) - {ceco}
            if query_name and query_name != sort_key(store["store"]) and other_cecos:
                skipped["ambiguousStoreName"] += 1
                continue
            try:
                hire, birth = iso_date(cell(row, columns, "F_INGRESO")), iso_date(cell(row, columns, "F.NAC"))
            except ValueError:
                skipped["invalidRow"] += 1
                continue
            if not hire or not birth or date.fromisoformat(hire) > cutoff or date.fromisoformat(birth) > cutoff:
                skipped["invalidRow"] += 1
                continue
            shift, role = clean_text(cell(row, columns, "TURNO")), clean_text(cell(row, columns, "NOM_PUESTO"))
            if not shift or not role:
                skipped["invalidRow"] += 1
                continue
            if header_key("NOM_CCOSTO") in columns and sort_key(cell(row, columns, "NOM_CCOSTO")) != sort_key(store["store"]):
                name_mismatches += 1
            values = (name, hire, ceco, store["store"], shift, role, store["region"], store["dm"], f"--{birth[5:]}")
            if number in invalid_employees:
                skipped["conflictingEmployee"] += 1
                continue
            if number in candidates:
                if candidates[number] != values:
                    candidates.pop(number)
                    invalid_employees.add(number)
                    skipped["conflictingEmployee"] += 2
                    continue
                identical_rows += 1
                continue
            candidates[number] = values
        published = [list(values) for values in candidates.values()]
        if not source_rows or not published:
            raise ValueError("Query no contiene partners activos válidos")
    finally:
        workbook.close()
    navigation = build_navigation(published, RUNTIME_HEADERS)
    unique_dms = {(region["name"], dm["name"]) for region in navigation for dm in region["dms"]}
    unique_stores = {row[RUNTIME_HEADERS.index("CECO")] for row in published}
    meta = {
        "source": excel.name, "directory": directory.name, "sheet": matches[0],
        "version": f"{cutoff:%Y-%m-%d}-directorio-ceco", "generatedOn": date.today().isoformat(),
        "periodEnd": cutoff.isoformat(),
        "month": cutoff.month, "monthName": MONTHS[cutoff.month - 1], "year": cutoff.year,
        "sourceRows": source_rows, "publishedRows": len(published),
        "validated": True, "excludedRows": source_rows - len(published),
        "stores": len(unique_stores), "regions": len(navigation),
        "sourceSha256": _digest(excel), "directorySha256": _digest(directory),
    }
    audit = {
        "status": "ok", **meta, "headerRow": first, "columnsRead": list(EXPECTED_HEADERS),
        "columnsPublished": list(RUNTIME_HEADERS), "terminatedExcluded": terminated,
        "identicalEmployeeRowsRemoved": identical_rows, "queryStoreNameDifferences": name_mismatches,
        "skippedByReason": dict(sorted(skipped.items())),
        "directoryAudit": directory_audit,
        "uniqueRegions": len(navigation), "uniqueDM": len(unique_dms), "uniqueStores": len(unique_stores),
        "baselinePublishedRows": _baseline_count(baseline_path),
    }
    js = "// Generado por engine/partner_engine. No editar manualmente.\n"
    for key, value in (("META", meta), ("NAV", navigation), ("HEADERS", RUNTIME_HEADERS), ("ROWS", published)):
        js += f"window.PARTNER_{key}=" + json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False) + ";\n"
    _atomic_write(output, js)
    _atomic_write(audit_output, json.dumps(audit, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    return BuildResult(source_rows, len(published), cutoff.isoformat(), cutoff.month, cutoff.year,
                       len(navigation), len(unique_dms), len(unique_stores), str(output_path), str(audit_path))
