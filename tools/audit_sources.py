"""Preflight de Query y Directorio. No escribe datos de empleados."""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from engine.partner_engine.generator import (
    EXPECTED_HEADERS, _digest, ceco_from_ccosto, cell, column_positions,
    find_header_row, iso_date, load_directory, query_cutoff, sort_key,
)


def audit_sources(query: Path, directory: Path) -> dict:
    stores, conflicts, directory_audit = load_directory(directory)
    names = defaultdict(set)
    for ceco, record in stores.items():
        names[sort_key(record["store"])].add(ceco)
    workbook = load_workbook(query, read_only=True, data_only=True)
    try:
        matching = [name for name in workbook.sheetnames if sort_key(name) == "query"]
        if len(matching) != 1:
            raise ValueError("Se requiere una sola hoja Query")
        sheet = workbook[matching[0]]
        cutoff = query_cutoff(sheet)
        header_row, headers = find_header_row(sheet, EXPECTED_HEADERS)
        columns = column_positions(headers)
        totals = Counter()
        missing, not_open, ambiguous = Counter(), Counter(), Counter()
        mismatched_names = Counter()
        matched_cecos = set()
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(value is not None for value in row):
                continue
            totals["sourceRows"] += 1
            termination = iso_date(cell(row, columns, "F_BAJA"))
            if termination and termination <= cutoff.isoformat():
                totals["terminatedAtCutoff"] += 1
                continue
            totals["activeAtCutoff"] += 1
            ceco = ceco_from_ccosto(cell(row, columns, "CCOSTO"))
            if not ceco:
                totals["invalidCCOSTO"] += 1
                continue
            if ceco in conflicts:
                totals["conflictingDirectoryRows"] += 1
                continue
            record = stores.get(ceco)
            if not record:
                missing[ceco] += 1
                continue
            if sort_key(record["status"]) != "abierta":
                not_open[(ceco, record["status"])] += 1
                continue
            query_name = sort_key(cell(row, columns, "NOM_CCOSTO")) if "nom_ccosto" in columns else ""
            if query_name != sort_key(record["store"]):
                mismatched_names[ceco] += 1
                others = names.get(query_name, set()) - {ceco}
                if others:
                    ambiguous[(ceco, ",".join(sorted(others)))] += 1
                    continue
            totals["matchedOpenRows"] += 1
            matched_cecos.add(ceco)
        unresolved = sum(missing.values()) + sum(not_open.values()) + sum(ambiguous.values()) + totals["invalidCCOSTO"] + totals["conflictingDirectoryRows"]
        published = query.parent / "partners.js"
        published_matches = False
        if published.is_file():
            match = re.search(r"window\.PARTNER_META=(.*?);", published.read_text(encoding="utf-8"))
            if match:
                meta = json.loads(match.group(1))
                published_matches = (meta.get("sourceSha256") == _digest(query)
                                     and meta.get("directorySha256") == _digest(directory)
                                     and meta.get("periodEnd") == cutoff.isoformat())
        return {
            "cutoffM1": cutoff.isoformat(), "sourceSha256": _digest(query),
            "directorySha256": _digest(directory), "canPublish": unresolved == 0,
            "publishedSnapshotMatchesSources": published_matches,
            "counts": dict(totals), "matchedOpenStores": len(matched_cecos),
            "queryNameDiffersFromDirectoryRows": sum(mismatched_names.values()),
            "unresolvedRows": unresolved, "directory": directory_audit,
            "missingCeCos": [{"ceco": c, "rows": n} for c, n in sorted(missing.items())],
            "nonOpenCeCos": [{"ceco": c, "status": s, "rows": n} for (c, s), n in sorted(not_open.items())],
            "ambiguousNameCeCos": [{"queryCeCo": c, "nameMatchesCeCos": other.split(","), "rows": n}
                                   for (c, other), n in sorted(ambiguous.items())],
            "exampleJoin": [{"ccosto": "01938262" if c == "38262" else "CeCo " + c,
                             "ceco": c, **stores[c]} for c in ("38262", "38362") if c in stores],
        }
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="Audita cruces sin publicar datos personales")
    parser.add_argument("--query", type=Path, default=Path("data/Query.xlsx"))
    parser.add_argument("--directory", type=Path, default=Path("data/Directorio.xlsx"))
    parser.add_argument("--output", type=Path, default=Path("reports/validacion_corte.json"))
    args = parser.parse_args()
    report = audit_sources(args.query, args.directory)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"canPublish": report["canPublish"], "unresolvedRows": report["unresolvedRows"],
                      "output": str(args.output)}, ensure_ascii=False))
    return 0 if report["canPublish"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
